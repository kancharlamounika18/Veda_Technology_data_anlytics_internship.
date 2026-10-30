"""Top 10 Products by Sales - Superstore task.

Usage:
    python top10_products_analysis.py superstore.csv top10_products.xlsx

The script groups sales by Product Name, sorts descending, prints the Top 10,
and writes an Excel workbook containing the raw data, Top-10 table, and chart.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SOURCE_URL = "https://github.com/leonism/sample-superstore/blob/master/data/superstore.csv"


def build_workbook(data: pd.DataFrame, output_path: Path) -> pd.DataFrame:
    """Create the Excel deliverable and return the ranked Top-10 table."""
    required = {"Product Name", "Sales"}
    missing = required - set(data.columns)
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(sorted(missing))}")

    data = data.copy()
    data["Sales"] = pd.to_numeric(data["Sales"], errors="raise")
    top10 = (
        data.groupby("Product Name", as_index=False)["Sales"]
        .sum()
        .sort_values(["Sales", "Product Name"], ascending=[False, True])
        .head(10)
        .reset_index(drop=True)
    )
    top10.insert(0, "Rank", range(1, len(top10) + 1))

    wb = Workbook()
    summary = wb.active
    summary.title = "Top 10"
    raw = wb.create_sheet("Data")

    # Summary sheet.
    summary["A1"] = "Top 10 Products by Sales"
    summary["A2"] = "Dataset: Tableau Sample Superstore | Metric: total Sales | Tie-break: product name A-Z"
    summary["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    summary.merge_cells("A1:C1")
    summary.merge_cells("A2:C2")
    summary["A2"].font = Font(italic=True, color="666666")
    headers = ["Rank", "Product Name", "Total Sales"]
    for col, value in enumerate(headers, 1):
        cell = summary.cell(4, col, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center")

    # Product names are the ranked result; sales are live SUMIF formulas against Data.
    data_product_col = get_column_letter(data.columns.get_loc("Product Name") + 1)
    data_sales_col = get_column_letter(data.columns.get_loc("Sales") + 1)
    for i, row in top10.iterrows():
        excel_row = 5 + i
        summary.cell(excel_row, 1, int(row["Rank"]))
        summary.cell(excel_row, 2, row["Product Name"])
        summary.cell(excel_row, 3, f'=SUMIF(Data!${data_product_col}:${data_product_col},B{excel_row},Data!${data_sales_col}:${data_sales_col})')
        summary.cell(excel_row, 3).number_format = '$#,##0.00'

    summary["A16"] = "Source"
    summary["B16"] = SOURCE_URL
    summary["A16"].font = Font(bold=True)
    summary["B16"].font = Font(color="0563C1", underline="single")
    summary["B16"].hyperlink = SOURCE_URL
    summary["A18"] = "Note"
    summary["B18"] = "Sales are summed by Product Name. Ties are ordered alphabetically."
    summary["A18"].font = Font(bold=True)
    summary["B18"].alignment = Alignment(wrap_text=True)

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Top 10 Products by Total Sales"
    chart.y_axis.title = "Product"
    chart.x_axis.title = "Sales ($)"
    chart.height = 9
    chart.width = 18
    chart.add_data(Reference(summary, min_col=3, min_row=4, max_row=14), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=2, min_row=5, max_row=14))
    chart.legend = None
    summary.add_chart(chart, "E4")

    # Raw data sheet.
    for col, value in enumerate(data.columns, 1):
        cell = raw.cell(1, col, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(data.itertuples(index=False, name=None), 2):
        for c_idx, value in enumerate(row, 1):
            raw.cell(r_idx, c_idx, value.item() if hasattr(value, "item") else value)
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = f"A1:{get_column_letter(len(data.columns))}{len(data) + 1}"
    ref = f"A1:{get_column_letter(len(data.columns))}{len(data) + 1}"
    table = Table(displayName="SuperstoreData", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    raw.add_table(table)

    for ws in (summary, raw):
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0 if ws.title == "Data" else 1
        ws.page_setup.orientation = "landscape"
    summary.column_dimensions["A"].width = 10
    summary.column_dimensions["B"].width = 55
    summary.column_dimensions["C"].width = 18
    summary.column_dimensions["D"].width = 3
    summary.column_dimensions["E"].width = 15
    for idx, col in enumerate(data.columns, 1):
        width = min(max(len(str(col)) + 2, 12), 32)
        raw.column_dimensions[get_column_letter(idx)].width = width
    for cell in raw[1]:
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return top10


def main() -> None:
    parser = argparse.ArgumentParser(description="Find and report the Top 10 Superstore products by sales.")
    parser.add_argument("input_csv", type=Path, help="Path to the Superstore CSV file")
    parser.add_argument("output_xlsx", type=Path, nargs="?", default=Path("top10_products.xlsx"), help="Output Excel path")
    args = parser.parse_args()
    df = pd.read_csv(args.input_csv)
    top10 = build_workbook(df, args.output_xlsx)
    print("Top 10 products by total sales:")
    print(top10.to_string(index=False))
    print(f"\nWorkbook written to: {args.output_xlsx.resolve()}")


if __name__ == "__main__":
    main()
