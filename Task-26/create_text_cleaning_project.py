from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path('/home/ubuntu/Text-Cleaning-Basics')
ROOT.mkdir(parents=True, exist_ok=True)

# Intentionally messy customer data for text-cleaning practice.
raw = pd.DataFrame([
    ['C001', '  aarav sharma ', ' AARAV.SHARMA@EMAIL.COM ', ' hyderabad ', ' 98765 43210 '],
    ['C002', 'MEERA   PATEL', 'meera.patel @email.com', 'MUMBAI', '9876543211'],
    ['C003', ' rohan kumar', 'ROHAN.KUMAR@EMAIL.COM', ' bengaluru', '  91234-56789'],
    ['C004', 'ANANYA singh ', ' ananya.singh@email.com ', 'Delhi ', '99887 76655'],
    ['C005', '  vikram rao', 'vikram.rao@EMAIL.COM', 'Chennai', '  9000011111 '],
    ['C006', 'diya nair  ', 'DIYA.NAIR @email.com', ' kochi', '88999-00011'],
    ['C007', 'KARAN  shah', 'karan.shah@email.com', 'Pune', ' 77777 88888'],
    ['C008', ' sneha iyer', ' sneha.iyer@email.com', 'BENGALURU ', '  96666 55555'],
    ['C009', 'rahul  reddy ', 'RAHUL.REDDY@EMAIL.COM ', ' hyderabad', '95555-44444'],
    ['C010', 'PRIYA  DESAI', 'priya.desai @EMAIL.COM', ' Ahmedabad ', '  94444 33333'],
    ['C011', '  aditya verma', 'ADITYA.VERMA@email.com', 'LUCKNOW', '93333 22221'],
    ['C012', 'kavya menon ', ' kavya.menon@email.com ', ' thiruvananthapuram ', '92222-11110'],
    ['C013', 'SAMEER khan', 'sameer.khan@EMAIL.COM', 'Jaipur ', '91111 00009'],
    ['C014', '  neha joshi  ', 'NEHA.JOSHI @email.com', ' surat', '  90000-12345 '],
    ['C015', 'AMIT  GUPTA', ' amit.gupta@email.com', 'NEW DELHI', '88888 99990'],
], columns=['customer_id','customer_name','email','city','phone'])
raw.to_csv(ROOT/'customer_dataset.csv', index=False)

# Cleaning rules: trim whitespace, normalize internal spaces, title-case names/cities,
# lowercase and remove spaces from email, and retain digits in phone numbers.
def clean_spaces(value):
    return ' '.join(str(value).strip().split())

def clean_name(value):
    return clean_spaces(value).title()

def clean_email(value):
    return clean_spaces(value).lower().replace(' ', '')

def clean_city(value):
    value = clean_spaces(value).title()
    city_map = {'Bengaluru':'Bengaluru', 'New Delhi':'New Delhi', 'Thiruvananthapuram':'Thiruvananthapuram'}
    return city_map.get(value, value)

def clean_phone(value):
    return ''.join(ch for ch in str(value) if ch.isdigit())

cleaned = raw.copy()
cleaned['customer_name'] = cleaned['customer_name'].map(clean_name)
cleaned['email'] = cleaned['email'].map(clean_email)
cleaned['city'] = cleaned['city'].map(clean_city)
cleaned['phone'] = cleaned['phone'].map(clean_phone)
cleaned.to_csv(ROOT/'cleaned_customer_dataset.csv', index=False)

# Write the requested reusable Python script.
script = '''"""Text Cleaning Basics — Veda Technology Data Analytics Internship.

Run from this project directory:
    python text_cleaning.py
"""
from pathlib import Path
import pandas as pd

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "customer_dataset.csv"
OUTPUT_FILE = BASE_DIR / "cleaned_customer_dataset.csv"


def clean_spaces(value):
    """Remove leading/trailing whitespace and collapse repeated spaces."""
    return " ".join(str(value).strip().split())


def clean_name(value):
    return clean_spaces(value).title()


def clean_email(value):
    return clean_spaces(value).lower().replace(" ", "")


def clean_city(value):
    return clean_spaces(value).title()


def clean_phone(value):
    return "".join(character for character in str(value) if character.isdigit())


def clean_customer_data(input_file=INPUT_FILE, output_file=OUTPUT_FILE):
    data = pd.read_csv(input_file, dtype=str).fillna("")
    data["customer_name"] = data["customer_name"].map(clean_name)
    data["email"] = data["email"].map(clean_email)
    data["city"] = data["city"].map(clean_city)
    data["phone"] = data["phone"].map(clean_phone)
    data.to_csv(output_file, index=False)
    return data


if __name__ == "__main__":
    cleaned_data = clean_customer_data()
    print(f"Cleaned {len(cleaned_data)} customer records.")
    print(f"Saved cleaned file to: {OUTPUT_FILE}")
'''
(ROOT/'text_cleaning.py').write_text(script, encoding='utf-8')

# Excel before/after workbook.
wb = load_workbook(ROOT/'before_after_examples.xlsx') if (ROOT/'before_after_examples.xlsx').exists() else None
from openpyxl import Workbook
wb = Workbook()
ws = wb.active; ws.title = 'Before_After'
rules = wb.create_sheet('Cleaning_Rules')
summary = wb.create_sheet('Summary')
navy='17365D'; blue='2F75B5'; light='D9EAF7'; pale='F4F8FB'; green='E2F0D9'; orange='FCE4D6'; white='FFFFFF'; gray='666666'; thin=Side(style='thin',color='D9E2F3')

def header(ws, title, subtitle, end_col):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=end_col)
    c=ws.cell(1,1,title); c.font=Font(size=16,bold=True,color=white); c.fill=PatternFill('solid',fgColor=navy); c.alignment=Alignment(horizontal='left')
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=end_col)
    c=ws.cell(2,1,subtitle); c.font=Font(size=10,italic=True,color=gray); c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=28

def style_header(ws,row,cols):
    for col in range(1,cols+1):
        c=ws.cell(row,col); c.font=Font(bold=True,color=white); c.fill=PatternFill('solid',fgColor=blue); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=Border(bottom=thin)
    ws.row_dimensions[row].height=24

def widths(ws, count):
    for col in range(1,count+1):
        letter=get_column_letter(col); maxlen=max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=10)
        ws.column_dimensions[letter].width=min(max(maxlen+2,14),38)

header(ws,'Text Cleaning — Before and After Examples','Raw values compared with standardized values.',10)
cols=['customer_id','raw_customer_name','clean_customer_name','raw_email','clean_email','raw_city','clean_city','raw_phone','clean_phone']
for j,c in enumerate(cols,1): ws.cell(4,j,c.replace('_',' ').title())
style_header(ws,4,len(cols))
for i,(before,after) in enumerate(zip(raw.itertuples(index=False),cleaned.itertuples(index=False)),5):
    vals=[before[0],before[1],after[1],before[2],after[2],before[3],after[3],before[4],after[4]]
    for j,v in enumerate(vals,1): ws.cell(i,j,v)
ref=f'A4:I{4+len(raw)}'; tab=Table(displayName='BeforeAfterExamples',ref=ref); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False); ws.add_table(tab)
ws.freeze_panes='A5'; widths(ws,9)
for row in range(5,5+len(raw)):
    for col in [2,4,6,8]: ws.cell(row,col).fill=PatternFill('solid',fgColor=orange)
    for col in [3,5,7,9]: ws.cell(row,col).fill=PatternFill('solid',fgColor=green)

header(rules,'Cleaning Rules','Functions applied to standardize the customer data.',4)
for j,c in enumerate(['Field','Cleaning Rule','Example Before','Example After'],1): rules.cell(4,j,c)
style_header(rules,4,4)
rule_rows=[('customer_name','Trim, collapse repeated spaces, title case','  aarav   sharma ','Aarav Sharma'),('email','Trim, lowercase, remove spaces',' AARAV.SHARMA @EMAIL.COM ','aarav.sharma@email.com'),('city','Trim, collapse spaces, title case',' BENGALURU ','Bengaluru'),('phone','Keep digits only',' 98765-43210 ','9876543210')]
for i,row in enumerate(rule_rows,5):
    for j,v in enumerate(row,1): rules.cell(i,j,v)
format_ref=f'A4:D{4+len(rule_rows)}'; tab=Table(displayName='CleaningRules',ref=format_ref); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False); rules.add_table(tab); widths(rules,4); rules.freeze_panes='A5'

header(summary,'Text Cleaning Summary','Quality checks and project details.',3)
summary_rows=[('Project','Text Cleaning Basics'),('Organization','Veda Technology'),('Tool','Python and Excel'),('Input records',len(raw)),('Output records',len(cleaned)),('Fields cleaned','customer_name, email, city, phone'),('Main objective','Standardize names, spaces, capitalization, emails, cities, and phone numbers')]
for i,(k,v) in enumerate(summary_rows,4):
    summary.cell(i,1,k).font=Font(bold=True,color=navy); summary.cell(i,1).fill=PatternFill('solid',fgColor=light); summary.cell(i,2,v).fill=PatternFill('solid',fgColor=pale); summary.cell(i,2).alignment=Alignment(wrap_text=True)
summary['A13']='Legend'; summary['A13'].font=Font(bold=True,color=navy)
summary['A14']='Orange cells'; summary['B14']='Before cleaning'; summary['A14'].fill=PatternFill('solid',fgColor=orange)
summary['A15']='Green cells'; summary['B15']='After cleaning'; summary['A15'].fill=PatternFill('solid',fgColor=green)
summary.column_dimensions['A'].width=24; summary.column_dimensions['B'].width=78
for sheet in wb.worksheets: sheet.sheet_view.showGridLines=False
wb.save(ROOT/'before_after_examples.xlsx')
print(ROOT)
print('raw rows',len(raw),'cleaned rows',len(cleaned))
